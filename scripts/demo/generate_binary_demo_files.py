#!/usr/bin/env python3
"""
Generate thesis-ready DOCX and PDF demo files under datasets/demo/resource-files/.

Deterministic, offline, no network. Requires: python-docx, fpdf2 (see project .venv).

Usage:
  .venv/bin/python scripts/demo/generate_binary_demo_files.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from fpdf import FPDF

OUT = Path(__file__).resolve().parents[2] / "datasets" / "demo" / "resource-files"


def _docx_add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _docx_add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def write_docx() -> Path:
    """Web Development lab: accessible forms, keyboard navigation, responsive layout."""
    path = OUT / "web-development-accessibility-lab.docx"
    doc = Document()

    doc.add_heading("Web Development Lab: Accessible Forms and Navigation", 0)
    doc.add_paragraph(
        "Course context: Web Development Learning Unit 098 (demo catalog). "
        "This lab applies semantic HTML, accessible forms, keyboard navigation, "
        "focus management, and responsive layout checks to a small course project page."
    )

    doc.add_heading("Intended learning outcomes", level=1)
    _docx_add_bullets(
        doc,
        [
            "Structure pages with semantic HTML (header, nav, main, footer).",
            "Associate every form control with a visible, programmatic label.",
            "Complete a keyboard-only walkthrough with visible focus states.",
            "Write meaningful alt text and avoid layout-only images without description.",
            "Document manual accessibility checks for a non-technical stakeholder summary.",
        ],
    )

    doc.add_heading("Semantic HTML", level=1)
    doc.add_paragraph(
        "Use headings in logical order (h1 once per page, then h2/h3). Prefer native "
        "elements: button for actions, a for navigation links, nav for site menus. "
        "Avoid div-only buttons; screen readers lose role and state information."
    )
    _docx_add_bullets(
        doc,
        [
            "Landmark regions: one main, optional complementary aside.",
            "Lists for navigation groups; tables only for tabular data.",
            "Language attribute on html element for correct pronunciation.",
        ],
    )

    doc.add_heading("Accessible forms", level=1)
    doc.add_paragraph(
        "Every input, select, and textarea needs a label linked via for/id or implicit "
        "wrapping. Placeholder text is not a label. Group related radios with fieldset "
        "and legend. Announce errors with text, not color alone."
    )
    _docx_add_bullets(
        doc,
        [
            "Required fields: aria-required or visible required indicator plus text.",
            "Error summary at top of form on submit failure.",
            "Autocomplete attributes where they help users (name, email).",
        ],
    )

    doc.add_heading("Keyboard navigation", level=1)
    doc.add_paragraph(
        "Tab order should follow visual reading order. Custom widgets must expose "
        "keyboard operability (Enter/Space on buttons, arrows in menus where expected). "
        "Trap focus only inside modal dialogs and restore focus on close."
    )

    doc.add_heading("Focus states", level=1)
    doc.add_paragraph(
        "Never remove outline without replacing it. Focus indicators need 3:1 contrast "
        "against adjacent colors. Skip link to main content is the first focusable "
        "element on long pages."
    )

    doc.add_heading("Images and alt text", level=1)
    doc.add_paragraph(
        "Decorative images: alt=\"\" (empty). Informative images: concise alt describing "
        "purpose (chart trend, not filename). Complex figures may need longer description "
        "in adjacent text."
    )

    doc.add_heading("Responsive layout", level=1)
    doc.add_paragraph(
        "Test at 320px and 1280px widths. Avoid horizontal scroll from fixed widths. "
        "Touch targets at least 44px where mobile users submit forms. Responsive layout "
        "supports zoom up to 200% without loss of content."
    )

    doc.add_heading("Lab tasks", level=1)
    _docx_add_numbered(
        doc,
        [
            "Fix a registration form missing labels and fieldset on a radio group.",
            "Add skip link and verify first Tab lands on it.",
            "Run keyboard-only path: open menu, submit form, close modal.",
            "Resize to 320px; fix overflow on the primary CTA row.",
            "Write a 150-word UX note for instructors explaining two fixes you made.",
        ],
    )

    doc.add_heading("Manual testing checklist", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Check"
    hdr[1].text = "Pass"
    hdr[2].text = "Notes"
    rows = [
        ("Headings in order", "", ""),
        ("All inputs labeled", "", ""),
        ("Visible focus on Tab", "", ""),
        ("Screen reader reads button names", "", ""),
        ("No horizontal scroll at 320px", "", ""),
    ]
    for check, pass_col, notes in rows:
        row = table.add_row().cells
        row[0].text = check
        row[1].text = pass_col
        row[2].text = notes

    doc.add_heading("Reflection questions", level=1)
    _docx_add_bullets(
        doc,
        [
            "Which change helped screen reader users the most?",
            "What would you automate in CI (axe, lint rules)?",
            "How does semantic HTML improve maintainability for your team?",
        ],
    )

    doc.add_heading("Expected submission", level=1)
    doc.add_paragraph(
        "Submit: (1) link or zip of updated HTML/CSS, (2) completed checklist table, "
        "(3) short screen recording or timestamped notes of keyboard walkthrough. "
        "Keywords: semantic HTML, accessibility, labels, keyboard navigation, "
        "screen reader, responsive layout, forms, UX."
    )

    doc.save(path)
    return path


class ThesisPdf(FPDF):
    """FPDF with helpers; Helvetica is ASCII-safe."""

    def section_title(self, title: str) -> None:
        self.set_font("Helvetica", "B", 12)
        self.multi_cell(self.epw, 7, title)
        self.ln(1)
        self.set_font("Helvetica", size=10)

    def body(self, text: str) -> None:
        self.multi_cell(self.epw, 5, text)
        self.ln(2)


def write_pdf() -> Path:
    """AI worksheet: classification metrics, confusion matrix, threshold trade-offs."""
    path = OUT / "ai-classification-metrics-worksheet.pdf"
    pdf = ThesisPdf()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(pdf.epw, 8, "AI Worksheet: Classification Metrics and Error Analysis")
    pdf.ln(2)
    pdf.set_font("Helvetica", size=10)

    pdf.body(
        "Context: In the Educational Companion System, supervised models and rules may "
        "classify learner needs (e.g., remediation vs enrichment). This worksheet "
        "practices model evaluation using a confusion matrix, precision, recall, F1-score, "
        "and threshold trade-offs. Original demo content for thesis ingestion (selectable text)."
    )

    pdf.section_title("Intended learning outcomes")
    pdf.body(
        "1) Build and interpret a binary confusion matrix. "
        "2) Compute accuracy, precision, recall, and F1. "
        "3) Explain when accuracy is misleading. "
        "4) Relate threshold changes to false positives and false negatives. "
        "5) Plan a short error analysis for one misclassified case."
    )

    pdf.section_title("Confusion matrix (binary classification)")
    pdf.body(
        "Labels: Positive = learner needs intervention. Negative = on track.\n"
        "TP = predicted positive, actually positive.\n"
        "FP = predicted positive, actually negative (false alarm).\n"
        "FN = predicted negative, actually positive (missed case).\n"
        "TN = predicted negative, actually negative."
    )

    pdf.section_title("Worked example")
    pdf.body(
        "Counts: TP=42, FN=8, FP=5, TN=145 (total N=200).\n"
        "Accuracy = (TP+TN)/N = 187/200 = 0.935.\n"
        "Precision = TP/(TP+FP) = 42/47 = 0.894 (when we alert, we are usually right).\n"
        "Recall = TP/(TP+FN) = 42/50 = 0.840 (we catch 84% of true intervention cases).\n"
        "F1 = 2*P*R/(P+R) = 2*0.894*0.840/(0.894+0.840) = 0.866."
    )

    pdf.section_title("Interpretation")
    pdf.body(
        "High precision protects learners from unnecessary nudges. High recall protects "
        "at-risk learners from being overlooked. For educational recommendations, false "
        "negatives may be costlier than false positives depending on policy."
    )

    pdf.add_page()
    pdf.section_title("Threshold trade-off")
    pdf.body(
        "Lowering the decision threshold usually increases recall and FP. Raising it "
        "increases precision and FN. Plot precision-recall for your dev set before "
        "choosing an operating point. Document the chosen threshold in the model card."
    )

    pdf.section_title("Error analysis")
    pdf.body(
        "Pick one FP: why did the model think intervention was needed? Missing feature, "
        "stale activity log, or ambiguous label? Pick one FN: was activity not ingested "
        "yet? Propose one feature or rule change and how you would validate it."
    )

    pdf.section_title("Practice tasks")
    pdf.body(
        "1) If FN cost is 3x FP cost, would you lower or raise threshold? Justify.\n"
        "2) Recalculate metrics if TP=48 and FN=2 with same FP,TN.\n"
        "3) When is accuracy misleading? Give class imbalance example.\n"
        "4) List two keywords you expect in an extracted summary: confusion matrix, F1-score."
    )

    pdf.section_title("Reflection questions")
    pdf.body(
        "How should access-aware filtering interact with evaluation (only labeled accessible items)? "
        "Why must ingestion summaries stay scoped to the parent resource visibility? "
        "How could richer PDF text improve semantic matching for AI units without exposing full files?"
    )

    pdf.section_title("Keywords")
    pdf.body(
        "classification metrics, confusion matrix, precision, recall, F1-score, error analysis, "
        "threshold, supervised learning, model evaluation."
    )

    pdf.output(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    docx_path = write_docx()
    pdf_path = write_pdf()
    print(f"Wrote {docx_path}")
    print(f"Wrote {pdf_path}")
    print("Done. Commit generated .docx/.pdf or regenerate before demos.")


if __name__ == "__main__":
    main()
